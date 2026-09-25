"""Автоматичне читання багатoаркушевих книг архівних описів."""
from __future__ import annotations

import fnmatch
import re
from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .models import Dataset, Description, Issue, Record
from .text_matching import clean_cell, detect_title_language, normalize_case_id, normalize_text


YEAR_RE = re.compile(r"(?<!\d)((?:17|18|19|20)\d{2})(?!\d)")
YEAR_HEADING_RE = re.compile(
    r"^\s*(?:17|18|19|20)\d{2}(?:\s*(?:[,;/]|-|–|—|і|та)\s*(?:17|18|19|20)\d{2})*"
    r"\s*(?:рік|роки|років|рр?\.?|год|годы|гг?\.?)?\s*$",
    re.I,
)
WITHDRAWN_RE = re.compile(r"^(?:в\s*и\s*б\s*у\s*л\s*[аио]|в\s*ы\s*б\s*ы\s*л\s*[аои])", re.I)
INVENTORY_RE = re.compile(
    r"^(?:(?:архівний опис|архивная опись|недействующая опись|недіючий опис)(?:\s|$|[.,])"
    r"|(?:опис|опись)\s*[№#]\s*\d+\s+(?:фонду|фонда)\b)",
    re.I,
)
PERSONAL_SECTION_RE = re.compile(
    r"^(?:особов[іи]\s+справ[и]|особист[іи]\s+справ[и]|личн(?:ые|ое)\s+дел[ао]"
    r"|(?:список|списки)\s+(?:учнів|учениць|студентів|службовців|працівників|особового\s+складу)"
    r"|(?:список|списки)\s+(?:учеников|учениц|студентов|служащих|работников|личного\s+состава))\b",
    re.I,
)
LETTER_HEADING_RE = re.compile(r'^[«"“]?\s*[А-ЯЁІЇЄҐA-Z]\s*[»"”]?$', re.I)
HEADER_MARKERS = (
    "заголовок справ",
    "крайні дати",
    "кількість аркуш",
    "примітк",
    "№ з/п",
    "№\\nз",
    "№ справ",
)


def safe_load_yaml(path: Path, yaml_module) -> dict[str, Any]:
    with path.open("r", encoding="utf-8") as stream:
        data = yaml_module.safe_load(stream)
    if not isinstance(data, dict):
        raise ValueError(f"Корінь YAML має бути словником: {path}")
    return data


def discover_workbooks(input_dir: Path) -> list[Path]:
    files = sorted(
        path for path in input_dir.iterdir()
        if path.is_file() and path.suffix.casefold() == ".xlsx" and not path.name.startswith("~$")
    )
    if len(files) < 2:
        raise ValueError(
            f"Для порівняння потрібно щонайменше два *.xlsx у {input_dir}. Знайдено: {len(files)}"
        )
    return files


def _slug(value: str) -> str:
    value = normalize_text(value)
    value = re.sub(r"[^0-9a-zа-яіїєґ]+", "_", value, flags=re.I).strip("_")
    return value or "dataset"


def _configured_dataset(path: Path, entries: list[dict[str, Any]]) -> Dataset | None:
    for entry in entries:
        if fnmatch.fnmatch(path.name.casefold(), str(entry.get("pattern", "")).casefold()):
            label = clean_cell(entry.get("label")) or path.stem
            return Dataset(
                id=clean_cell(entry.get("id")) or _slug(path.stem),
                label=label,
                short_label=clean_cell(entry.get("short_label")) or label,
                reference=clean_cell(entry.get("reference")),
                path=path,
                configured=True,
                minimum_year=int(entry["minimum_year"]) if entry.get("minimum_year") is not None else None,
                maximum_year=int(entry["maximum_year"]) if entry.get("maximum_year") is not None else None,
            )
    return None


def make_dataset(path: Path, entries: list[dict[str, Any]]) -> Dataset:
    configured = _configured_dataset(path, entries)
    if configured:
        return configured
    return Dataset(
        id=_slug(path.stem),
        label=path.stem.replace("_", " "),
        short_label=path.stem.replace("_", " "),
        reference="",
        path=path,
        configured=False,
    )


def _metadata(sheet) -> dict[str, str]:
    result: dict[str, str] = {}
    keys = {
        "архів": "archive",
        "архив": "archive",
        "фонд": "fond",
        "опис": "inventory",
        "опись": "inventory",
        "мова": "language",
        "язык": "language",
    }
    known_rows = sheet.max_row if isinstance(sheet.max_row, int) else 15
    for row in range(1, min(known_rows, 15) + 1):
        key = normalize_text(clean_cell(sheet.cell(row, 1).value)).strip(" :.\n\t")
        if key in keys:
            result[keys[key]] = clean_cell(sheet.cell(row, 2).value)
    return result


def _is_header_row(values: list[str]) -> bool:
    joined = " | ".join(normalize_text(value).replace("\n", " ") for value in values if value)
    return any(marker.replace("\n", " ") in joined for marker in HEADER_MARKERS)


def _find_header_row(sheet) -> int:
    scores: list[tuple[int, int]] = []
    known_rows = sheet.max_row if isinstance(sheet.max_row, int) else 25
    for row in range(1, min(known_rows, 25) + 1):
        values = [clean_cell(sheet.cell(row, column).value) for column in range(1, 6)]
        joined = " | ".join(normalize_text(value).replace("\n", " ") for value in values if value)
        score = sum(marker.replace("\n", " ") in joined for marker in HEADER_MARKERS)
        if score:
            scores.append((score, row))
    return max(scores, default=(0, 1), key=lambda item: (item[0], -item[1]))[1]


def _looks_case_id(value: str) -> bool:
    normalized = normalize_case_id(value)
    return bool(normalized and re.match(r"^\d", normalized) and len(normalized) <= 30)


def _looks_year_heading(value: str) -> bool:
    return bool(value and YEAR_HEADING_RE.fullmatch(normalize_text(value).replace("\n", " ").strip()))


def _parse_years(value: str, section_year: int | None, minimum: int, maximum: int) -> tuple[int | None, int | None]:
    years = [int(match) for match in YEAR_RE.findall(value)]
    years = [year for year in years if minimum <= year <= maximum]
    if not years and section_year is not None:
        return section_year, section_year
    return (min(years), max(years)) if years else (None, None)


def _parse_pages(value: str) -> int | None:
    numbers = [int(item) for item in re.findall(r"(?<!\d)\d{1,5}(?!\d)", value)]
    return numbers[0] if len(numbers) == 1 else None


def _language(title: str, configured: str) -> str:
    detected = detect_title_language(title)
    if detected in {"uk", "ru"}:
        return detected
    normalized = normalize_text(configured)
    if normalized.startswith("ru") or "рос" in normalized:
        return "ru"
    return "uk"


def _fill_missing(record: Record, dates: str, pages: str, notes: str) -> bool:
    changed = False
    if dates and not record.dates_raw:
        record.dates_raw = dates
        changed = True
    if pages and not record.pages_raw:
        record.pages_raw = pages
        changed = True
    if notes and not record.notes:
        record.notes = notes
        changed = True
    return changed


def _make_record(
    dataset: Dataset,
    description: Description,
    row: int,
    case_id: str,
    title: str,
    dates: str,
    pages: str,
    notes: str,
    section_label: str,
    section_year: int | None,
    minimum_year: int,
    maximum_year: int,
    thematic_section: str = "",
) -> Record:
    normalized_title = normalize_text(title).strip()
    if WITHDRAWN_RE.match(normalized_title):
        status = "withdrawn"
    elif INVENTORY_RE.match(normalized_title):
        status = "inventory"
    else:
        status = "case"
    start_year, end_year = _parse_years(dates, section_year, minimum_year, maximum_year)
    return Record(
        dataset_id=dataset.id,
        dataset_label=dataset.label,
        file_name=dataset.path.name,
        sheet_name=description.sheet_name,
        excel_row=row,
        archive=description.archive,
        fond=description.fond,
        inventory=description.inventory,
        case_id=normalize_case_id(case_id),
        title=title,
        dates_raw=dates,
        pages_raw=pages,
        notes=notes,
        status=status,
        section_label=section_label,
        section_year=section_year,
        thematic_section=thematic_section,
        start_year=start_year,
        end_year=end_year,
        pages=_parse_pages(pages),
        language=_language(title, description.language),
    )


def read_sheet(
    dataset: Dataset,
    sheet,
    chronology: dict[str, Any],
) -> tuple[Description, list[Record], list[Issue]]:
    metadata = _metadata(sheet)
    archive = metadata.get("archive", "")
    fond = metadata.get("fond", "")
    inventory = metadata.get("inventory", "")
    language = metadata.get("language", "")
    header_row = _find_header_row(sheet)
    reference = ", ".join(
        part for part in (
            archive,
            f"ф. {fond}" if fond else "",
            f"оп. {inventory}" if inventory else "",
        ) if part
    )
    description = Description(
        dataset_id=dataset.id,
        file_name=dataset.path.name,
        sheet_name=sheet.title,
        archive=archive,
        fond=fond,
        inventory=inventory,
        language=language,
        header_row=header_row,
        reference=reference or sheet.title,
    )
    issues: list[Issue] = []
    records: list[Record] = []
    minimum_year = dataset.minimum_year or int(chronology.get("minimum_year", 1700))
    maximum_year = dataset.maximum_year or int(chronology.get("maximum_year", 2026))
    section_label = ""
    section_year: int | None = None
    thematic_section = ""
    pending_case_id = ""
    pending_row: int | None = None
    carry_dates = ""
    carry_pages = ""
    carry_notes = ""

    for row, raw_values in enumerate(
        sheet.iter_rows(min_row=header_row + 1, max_col=5, values_only=True),
        start=header_row + 1,
    ):
        values = [clean_cell(value) for value in raw_values]
        case_id, title, dates, pages, notes = values
        if not any(values):
            continue
        if _looks_year_heading(title):
            years = [int(year) for year in YEAR_RE.findall(title)]
            section_label = title
            section_year = years[0] if len(years) == 1 else None

            if records and _fill_missing(records[-1], dates, pages, notes):
                records[-1].start_year, records[-1].end_year = _parse_years(
                    records[-1].dates_raw,
                    records[-1].section_year,
                    minimum_year,
                    maximum_year,
                )
                records[-1].pages = _parse_pages(records[-1].pages_raw)
                issues.append(Issue(
                    "INFO", dataset.path.name, sheet.title, row,
                    "Зміщений рядок", " | ".join(value for value in (dates, pages, notes) if value),
                    f"Дати/аркуші перенесено до попередньої справи {records[-1].case_id}.",
                ))
            else:
                carry_dates, carry_pages, carry_notes = dates, pages, notes

            if _looks_case_id(case_id):
                pending_case_id = case_id
                pending_row = row
            continue

        if _is_header_row(values):
            continue

        if not case_id and title and pending_case_id:
            record = _make_record(
                dataset, description, pending_row or row, pending_case_id, title,
                dates or carry_dates, pages or carry_pages, notes or carry_notes,
                section_label, section_year, minimum_year, maximum_year,
                thematic_section,
            )
            records.append(record)
            issues.append(Issue(
                "INFO", dataset.path.name, sheet.title, row,
                "Зміщений рядок", pending_case_id,
                f"Заголовок об'єднано з номером справи з рядка {pending_row}.",
            ))
            pending_case_id = ""
            pending_row = None
            carry_dates = carry_pages = carry_notes = ""
            continue

        if not case_id and title:
            if carry_dates or carry_pages or carry_notes:
                synthetic = f"без-номера-{row}"
                records.append(_make_record(
                    dataset, description, row, synthetic, title,
                    dates or carry_dates, pages or carry_pages, notes or carry_notes,
                    section_label, section_year, minimum_year, maximum_year,
                    thematic_section,
                ))
                issues.append(Issue(
                    "WARNING", dataset.path.name, sheet.title, row,
                    "Номер справи", "", "Створено службовий ідентифікатор для заголовка без номера.",
                ))
                carry_dates = carry_pages = carry_notes = ""
            else:
                section_label = title
                years = [int(year) for year in YEAR_RE.findall(title)]
                section_year = years[0] if len(years) == 1 else None
                if PERSONAL_SECTION_RE.match(title):
                    thematic_section = title
                elif not LETTER_HEADING_RE.fullmatch(title):
                    thematic_section = ""
            continue

        if _looks_case_id(case_id):
            record = _make_record(
                dataset, description, row, case_id, title, dates, pages, notes,
                section_label, section_year, minimum_year, maximum_year,
                thematic_section,
            )
            records.append(record)
            if not title:
                issues.append(Issue(
                    "WARNING", dataset.path.name, sheet.title, row,
                    "Заголовок справи", "", "Номер справи є, але заголовок відсутній.",
                ))
            continue

        issues.append(Issue(
            "WARNING", dataset.path.name, sheet.title, row,
            "Структура рядка", " | ".join(value for value in values if value),
            "Рядок не вдалося однозначно віднести до справи або службового заголовка.",
        ))

    if pending_case_id:
        issues.append(Issue(
            "WARNING", dataset.path.name, sheet.title, pending_row,
            "Номер справи", pending_case_id,
            "Після номера, зміщеного в рядок річного заголовка, не знайдено заголовок справи.",
        ))

    for record in records:
        years = [int(year) for year in YEAR_RE.findall(record.dates_raw)]
        outside = sorted({year for year in years if year < minimum_year or year > maximum_year})
        if outside:
            issues.append(Issue(
                "WARNING", dataset.path.name, sheet.title, record.excel_row,
                "Крайні дати", record.dates_raw,
                "Рік виходить за очікувані межі "
                f"{minimum_year}–{maximum_year}: " + ", ".join(map(str, outside)) + ".",
            ))
        if record.pages_raw and record.pages is None:
            issues.append(Issue(
                "WARNING", dataset.path.name, sheet.title, record.excel_row,
                "Кількість аркушів", record.pages_raw,
                "Не вдалося однозначно визначити одне числове значення.",
            ))

    occurrences: dict[str, list[Record]] = defaultdict(list)
    for record in records:
        if record.case_id:
            occurrences[record.case_id].append(record)
    for case_id, duplicated in occurrences.items():
        if len(duplicated) > 1:
            issues.append(Issue(
                "WARNING", dataset.path.name, sheet.title, duplicated[0].excel_row,
                "Номер справи", case_id,
                "Номер повторюється у рядках " + ", ".join(str(item.excel_row) for item in duplicated) + ".",
            ))
    return description, records, issues


def load_all(
    input_dir: Path,
    dataset_config: dict[str, Any],
    chronology: dict[str, Any],
) -> tuple[list[Dataset], list[Description], list[Record], list[Issue]]:
    files = discover_workbooks(input_dir)
    entries = dataset_config.get("datasets", [])
    if not isinstance(entries, list):
        raise ValueError("config/datasets.yaml: datasets має бути списком")

    datasets: list[Dataset] = []
    descriptions: list[Description] = []
    records: list[Record] = []
    issues: list[Issue] = []
    used_ids: set[str] = set()

    for path in files:
        dataset = make_dataset(path, entries)
        if dataset.id in used_ids:
            raise ValueError(f"Повторюється id логічного масиву: {dataset.id}")
        used_ids.add(dataset.id)
        workbook = load_workbook(path, read_only=True, data_only=True)
        dataset_descriptions: list[Description] = []
        for sheet in workbook.worksheets:
            description, sheet_records, sheet_issues = read_sheet(dataset, sheet, chronology)
            dataset_descriptions.append(description)
            descriptions.append(description)
            records.extend(sheet_records)
            issues.extend(sheet_issues)
        workbook.close()

        if not dataset.reference:
            refs = list(dict.fromkeys(item.reference for item in dataset_descriptions if item.reference))
            dataset.reference = "; ".join(refs)
        datasets.append(dataset)
        if not dataset.configured:
            issues.append(Issue(
                "INFO", path.name, "", None, "Конфігурація масиву", dataset.label,
                "Файл не має окремого запису в datasets.yaml; назву утворено з імені файла.",
            ))
    return datasets, descriptions, records, issues
