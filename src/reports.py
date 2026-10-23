'CSV, Excel, HTML and charts for comparative analysis.'
from __future__ import annotations

import csv
import hashlib
import html
import json
import platform
import re
from importlib.metadata import version
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from . import __version__
from .localization import Visualization
from .models import Category, Dataset, Description, Issue, Record
from .comparative_profiles import REPORT_COLUMNS
from .comparative_profiles import DOCUMENT_TYPE_RULES_VERSION


NAVY = "1F4E78"
PALE = "F6F8FA"
WHITE = "FFFFFF"
TEXT = "202020"
WARNING = "FFF2CC"
ERROR = "F4CCCC"


def _write_csv(path: Path, rows: list[dict[str, Any]], fields: list[str] | None = None) -> None:
    if not rows and not fields:
        path.write_text("", encoding="utf-8-sig")
        return
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields or list(rows[0]), delimiter=";")
        writer.writeheader()
        writer.writerows(rows)


def record_rows(records: list[Record], visual: Visualization | None = None) -> list[dict[str, Any]]:
    rows = [{
        "record_uid": record.uid,
        "dataset_id": record.dataset_id,
        "dataset": record.dataset_label,
        "file_name": record.file_name,
        "sheet_name": record.sheet_name,
        "excel_row": record.excel_row,
        "archive": record.archive,
        "fond": record.fond,
        "inventory": record.inventory,
        "case_id": record.case_id,
        "title": record.title,
        "thematic_section": record.thematic_section,
        "dates_raw": record.dates_raw,
        "start_year": record.start_year,
        "end_year": record.end_year,
        "pages": record.pages,
        "notes": record.notes,
        "status": record.status,
        "language": record.language,
        "categories": " | ".join(record.category_labels),
        "category_ids": " | ".join(record.categories),
        "category_evidence": " | ".join(
            f"{label}: {', '.join(record.evidence.get(category_id, []))}"
            for category_id, label in zip(record.categories, record.category_labels)
        ),
        "context_categories": " | ".join(record.context_category_labels),
        "review_flags": " | ".join(record.review_flags),
    } for record in records]
    if visual is not None:
        for row, record in zip(rows, records):
            labels = [visual.label("categories", key, label)
                      for key, label in zip(record.categories, record.category_labels)]
            row["categories"] = " | ".join(labels)
            row["context_categories"] = " | ".join(
                visual.label("categories", key, label)
                for key, label in zip(record.context_categories, record.context_category_labels))
            row["category_evidence"] = " | ".join(
                f"{label}: " + ", ".join(
                    visual.text(item.partition(": ")[0]) + ": " + item.partition(": ")[2]
                    for item in record.evidence.get(key, []))
                for key, label in zip(record.categories, labels))
    return rows


def issue_rows(issues: list[Issue]) -> list[dict[str, Any]]:
    return [{
        "level": issue.level,
        "file_name": issue.file_name,
        "sheet_name": issue.sheet_name,
        "row": issue.row,
        "field": issue.field,
        "value": issue.value,
        "message": issue.message,
    } for issue in issues]


def write_csv_reports(
    tables_dir: Path,
    analysis: dict[str, Any],
    records: list[Record],
    issues: list[Issue],
) -> None:
    tables_dir.mkdir(parents=True, exist_ok=True)
    mapping = {
        "overview.csv": analysis["overview"],
        "classification_coverage.csv": analysis["classification_coverage"],
        "topic_unclassified_cases.csv": analysis["topic_unclassified"],
        "unclassified_cases.csv": analysis["unclassified"],
        "descriptions.csv": analysis["descriptions"],
        "chronology.csv": analysis["chronology"],
        "chronology_by_description.csv": analysis["chronology_by_description"],
        "categories.csv": analysis["categories"],
        "categories_by_description.csv": analysis["categories_by_description"],
        "classification_by_description.csv": analysis["classification_by_description"],
        "vocabulary.csv": analysis["vocabulary"],
        "common_vocabulary.csv": analysis["common_vocabulary"],
        "similarities.csv": analysis["similarities"],
        "records.csv": record_rows(records),
        "issues.csv": issue_rows(issues),
    }
    for filename, rows in mapping.items():
        _write_csv(tables_dir / filename, rows)
    for key, columns in REPORT_COLUMNS.items():
        _write_csv(tables_dir / f"{key}.csv", analysis[key], [key for key, _ in columns])


def _title(sheet, text: str, subtitle: str = "") -> int:
    sheet.sheet_view.showGridLines = False
    sheet.cell(2, 1, text)
    sheet.cell(2, 1).font = Font(name="Arial", size=14, bold=True, color=TEXT)
    sheet.cell(3, 1).fill = PatternFill("solid", fgColor=NAVY)
    sheet.row_dimensions[3].height = 3
    row = 5
    if subtitle:
        sheet.cell(4, 1, subtitle)
        sheet.cell(4, 1).font = Font(name="Arial", size=10, italic=True, color="666666")
        row = 6
    return row


def _tabular_sheet(
    workbook: Workbook,
    title: str,
    heading: str,
    columns: list[tuple[str, str]],
    rows: list[dict[str, Any]],
    subtitle: str = "",
    percent_keys: set[str] | None = None,
    decimal_keys: set[str] | None = None,
) -> None:
    sheet = workbook.create_sheet(title)
    start_row = _title(sheet, heading, subtitle)
    percent_keys = percent_keys or set()
    decimal_keys = decimal_keys or set()
    for column_index, (_, label) in enumerate(columns, start=1):
        sheet.cell(3, column_index).fill = PatternFill("solid", fgColor=NAVY)
        cell = sheet.cell(start_row, column_index, label)
        cell.font = Font(name="Arial", size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row_index, row in enumerate(rows, start=start_row + 1):
        for column_index, (key, _) in enumerate(columns, start=1):
            value = row.get(key)
            cell = sheet.cell(row_index, column_index, value)
            if isinstance(value, str):
                cell.data_type = "s"
            cell.font = Font(name="Arial", size=10, color=TEXT)
            cell.alignment = Alignment(vertical="top", wrap_text=key in {"title", "title_a", "title_b", "message", "reference"})
            if key in percent_keys and isinstance(value, (int, float)):
                cell.value = value / 100
                cell.number_format = "0.00%"
            elif key in decimal_keys and isinstance(value, (int, float)):
                cell.number_format = "0.0000"
        if row_index % 2 == 0:
            for column_index in range(1, len(columns) + 1):
                sheet.cell(row_index, column_index).fill = PatternFill("solid", fgColor=PALE)
    if rows:
        end_row = start_row + len(rows)
        table = Table(
            displayName="T_" + "".join(character for character in title if character.isalnum())[:20] + str(len(workbook.worksheets)),
            ref=f"A{start_row}:{get_column_letter(len(columns))}{end_row}",
        )
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        sheet.add_table(table)
    sheet.freeze_panes = f"A{start_row + 1}"
    sheet.auto_filter.ref = f"A{start_row}:{get_column_letter(len(columns))}{start_row + max(1, len(rows))}"
    sheet.row_dimensions[start_row].height = 32
    for column_index, (key, label) in enumerate(columns, start=1):
        sample = [str(row.get(key, "") or "") for row in rows[:400]]
        longest = max([len(label), *(min(len(value), 80) for value in sample)], default=len(label))
        width = min(max(longest + 2, 10), 55)
        if key in {"title", "title_a", "title_b", "message"}:
            width = 52
        sheet.column_dimensions[get_column_letter(column_index)].width = width


def _chronology_wide(datasets: list[Dataset], rows: list[dict[str, Any]], visual: Visualization) -> tuple[list[tuple[str, str]], list[dict[str, Any]]]:
    columns = [("year", 'Year')] + [(dataset.id, visual.dataset(dataset.id, dataset.short_label, short=True)) for dataset in datasets]
    return columns, rows


def _category_wide(
    datasets: list[Dataset],
    categories: list[Category],
    rows: list[dict[str, Any]],
    visual: Visualization,
) -> tuple[list[tuple[str, str]], list[dict[str, Any]], set[str]]:
    lookup = {(row["dataset_id"], row["category_id"]): row for row in rows}
    columns = [("category", 'Category')]
    percent_keys: set[str] = set()
    for dataset in datasets:
        columns.append((f"{dataset.id}_cases", f"{visual.dataset(dataset.id, dataset.short_label, short=True)}: {visual.text('Case files')}"))
        key = f"{dataset.id}_percent"
        columns.append((key, f"{visual.dataset(dataset.id, dataset.short_label, short=True)}: {visual.text('Share')}"))
        percent_keys.add(key)
    wide_rows: list[dict[str, Any]] = []
    for category in categories:
        item: dict[str, Any] = {"category": visual.label("categories", category.id, category.label)}
        for dataset in datasets:
            source = lookup[(dataset.id, category.id)]
            item[f"{dataset.id}_cases"] = source["cases"]
            item[f"{dataset.id}_percent"] = source["percent_of_titles"]
        wide_rows.append(item)
    return columns, wide_rows, percent_keys


def _common_wide(
    datasets: list[Dataset], rows: list[dict[str, Any]], visual: Visualization
) -> tuple[list[tuple[str, str]], list[dict[str, Any]], set[str]]:
    grouped: dict[int, dict[str, Any]] = {}
    percent_keys: set[str] = set()
    for row in rows:
        item = grouped.setdefault(row["rank"], {
            "rank": row["rank"],
            "term": row["term"],
            "average_share": row["average_share"],
        })
        item[f"{row['dataset_id']}_titles"] = row["titles"]
        item[f"{row['dataset_id']}_percent"] = row["percent_of_titles"]
    columns = [("rank", "№"), ("term", 'Term'), ("average_share", 'Mean share')]
    percent_keys.add("average_share")
    for dataset in datasets:
        columns.append((f"{dataset.id}_titles", f"{visual.dataset(dataset.id, dataset.short_label, short=True)}: {visual.text('Case titles')}"))
        key = f"{dataset.id}_percent"
        columns.append((key, f"{visual.dataset(dataset.id, dataset.short_label, short=True)}: {visual.text('Share')}"))
        percent_keys.add(key)
    return columns, list(grouped.values()), percent_keys


def write_workbook(
    path: Path,
    datasets: list[Dataset],
    descriptions: list[Description],
    records: list[Record],
    issues: list[Issue],
    categories: list[Category],
    analysis: dict[str, Any],
    visual: Visualization | None = None,
) -> None:
    visual = visual or Visualization()
    workbook = Workbook()
    workbook.remove(workbook.active)

    def tabulate(workbook, title, heading, columns, rows, subtitle="", **options):
        _tabular_sheet(workbook, visual.text(title), visual.text(heading),
                       [(key, visual.text(label)) for key, label in columns],
                       visual.rows(rows), visual.text(subtitle), **options)

    tabulate(
        workbook, 'Overview', 'Archival description comparison',
        [
            ("dataset", 'Dataset'), ("reference", 'Dataset composition'),
            ("descriptions", 'Inventories'), ("cases", 'Case files'),
            ("analyzable_titles", 'Analyzable titles'), ("withdrawn", 'Withdrawn'),
            ("start_year", 'Start year'), ("end_year", 'End year'),
            ("pages_total", 'Total leaves'), ("pages_mean", 'Mean leaves'),
            ("pages_median", 'Median leaves'), ("missing_titles", 'Missing title'),
            ("missing_dates", 'Missing dates'), ("missing_pages", 'Missing leaf count'),
            ("classified_titles", 'Subject classified'),
            ("classification_coverage_percent", 'Classification coverage'),
            ("context_only_titles", 'Context only'),
            ("context_only_percent", 'Context only, %'),
            ("unclassified_titles", 'Unclassified'),
            ("unclassified_percent", 'Unclassified, %'),
        ],
        analysis["overview"],
        subtitle='Each Excel workbook is one dataset; all its worksheets are aggregated.',
        percent_keys={
            "classification_coverage_percent", "context_only_percent",
            "unclassified_percent",
        },
    )
    tabulate(
        workbook, 'Classification', 'Thematic classification coverage',
        [
            ("dataset", 'Dataset'), ("reference", 'Dataset composition'),
            ("analyzable_titles", 'All analyzable titles (100%)'),
            ("subject_classified", 'Subject classified'),
            ("subject_classified_percent", 'Subject classified, %'),
            ("context_only", 'Context only'),
            ("context_only_percent", 'Context only, %'),
            ("unclassified", 'Unclassified'),
            ("unclassified_percent", 'Unclassified, %'),
        ],
        analysis["classification_coverage"],
        subtitle=(
            'Each dataset uses all its analyzable titles as 100%; the three exclusive states sum to 100%.'
        ),
        percent_keys={
            "subject_classified_percent", "context_only_percent",
            "unclassified_percent",
        },
    )
    tabulate(
        workbook, 'Without subject category', 'Case files without a subject category',
        [
            ("dataset", 'Dataset'), ("file_name", 'File'),
            ("sheet_name", 'Worksheet'), ("excel_row", 'Excel row'),
            ("archive", 'Archive'), ("fond", 'Fond'),
            ("inventory", 'Inventory'), ("case_id", 'Case number'),
            ("title", 'Title'), ("dates_raw", 'Inclusive dates'),
            ("classification_status", 'Status'),
            ("context_categories", 'Context categories'),
            ("review_flags", 'Review flags'),
        ],
        analysis["topic_unclassified"],
        subtitle=(
            'All case files without a subject category. Filter Status to distinguish context-only and unclassified titles.'
        ),
    )
    tabulate(
        workbook, 'Inventories', 'Inventory control totals',
        [
            ("dataset", 'Dataset'), ("file_name", 'File'), ("sheet_name", 'Worksheet'),
            ("archive", 'Archive'), ("fond", 'Fond'), ("inventory", 'Inventory'),
            ("cases", 'Case files'), ("analyzable_titles", 'Analyzable titles'),
            ("withdrawn", 'Withdrawn'), ("classified_titles", 'Classified'),
            ("missing_dates", 'Missing dates'), ("missing_pages", 'Missing leaf count'),
        ], analysis["descriptions"],
    )
    chronology_columns, chronology_data = _chronology_wide(datasets, analysis["chronology"], visual)
    tabulate(
        workbook, 'Chronology', 'Case files present by year',
        chronology_columns, chronology_data,
        subtitle='A case file spanning several years is counted in each year of that range.',
    )
    tabulate(
        workbook, 'Inventory chronology', 'Case files per inventory by year',
        [("dataset", 'Dataset'), ("sheet_name", 'Worksheet'), ("reference", 'Inventory'),
         ("year", 'Year'), ("cases", 'Case files')],
        analysis["chronology_by_description"],
        subtitle='Date ranges contribute to each year; inventories within a workbook remain separate.',
    )
    category_columns, category_data, category_percent = _category_wide(datasets, categories, analysis["categories"], visual)
    tabulate(
        workbook, 'Categories', 'Thematic structure of datasets',
        category_columns, category_data,
        subtitle=(
            'Each dataset uses its own analyzable titles as 100%. A title can have multiple categories, so shares can sum to more than 100%.'
        ),
        percent_keys=category_percent,
    )
    tabulate(
        workbook, 'Inventory categories', 'Categories by inventory',
        [("dataset", 'Dataset'), ("sheet_name", 'Worksheet'), ("reference", 'Inventory'),
         ("category", 'Category'), ("cases", 'Case files'),
         ("description_titles_total", 'Analyzable titles in inventory'),
         ("percent_of_description_titles", 'Inventory share'),
         ("dataset_titles_total", 'Case files in the corpus'),
         ("percent_of_dataset_titles", 'Dataset share')],
        analysis["categories_by_description"],
        subtitle='Dataset shares use all analyzable titles in that dataset; inventory shares use only titles in that worksheet.',
        percent_keys={"percent_of_description_titles", "percent_of_dataset_titles"},
    )
    tabulate(
        workbook, 'Inventory classification', 'Classification coverage by inventory',
        [("dataset", 'Dataset'), ("sheet_name", 'Worksheet'), ("reference", 'Inventory'),
         ("analyzable_titles", 'Analyzable titles in inventory'),
         ("subject_classified", 'Subject classified'),
         ("context_only", 'Context only'), ("unclassified", 'Unclassified'),
         ("subject_classified_percent", 'Classified, %'),
         ("context_only_percent", 'Context only, %'),
         ("unclassified_percent", 'Unclassified, %')],
        analysis["classification_by_description"],
        subtitle='Three exclusive states; shares use the analyzable titles in the corresponding inventory.',
        percent_keys={"subject_classified_percent", "context_only_percent", "unclassified_percent"},
    )
    tabulate(
        workbook, 'Distinctive vocabulary', 'Frequent and distinctive vocabulary',
        [
            ("dataset", 'Dataset'), ("type", 'List'), ("rank", 'Rank'),
            ("term", 'Term'), ("titles", 'Case titles'),
            ("percent_of_titles", 'Share of titles'), ("score", 'Distinctiveness score'),
        ], analysis["vocabulary"],
        percent_keys={"percent_of_titles"}, decimal_keys={"score"},
    )
    common_columns, common_data, common_percent = _common_wide(datasets, analysis["common_vocabulary"], visual)
    tabulate(
        workbook, 'Shared vocabulary', 'Vocabulary shared by all datasets',
        common_columns, common_data, percent_keys=common_percent,
    )
    tabulate(
        workbook, 'Similar case files', 'Lexically similar case titles across datasets',
        [
            ("similarity", 'Similarity'),
            ("dataset_a", 'Dataset A'), ("reference_a", 'Reference A'),
            ("case_a", 'Case A'), ("title_a", 'Title A'),
            ("dataset_b", 'Dataset B'), ("reference_b", 'Reference B'),
            ("case_b", 'Case B'), ("title_b", 'Title B'),
        ], analysis["similarities"], decimal_keys={"similarity"},
    )
    tabulate(
        workbook, 'Case files', 'Normalized case register',
        [
            ("dataset", 'Dataset'), ("sheet_name", 'Worksheet'), ("excel_row", 'Excel row'),
            ("archive", 'Archive'), ("fond", 'Fond'), ("inventory", 'Inventory'),
            ("case_id", 'Case number'), ("title", 'Title'), ("dates_raw", 'Inclusive dates'),
            ("thematic_section", 'Inventory subject heading'),
            ("start_year", 'Start year'), ("end_year", 'End year'),
            ("pages", 'Leaves'), ("status", 'Status'), ("language", 'Language'),
            ("categories", 'Categories'),
            ("category_evidence", 'Classification evidence'),
        ], record_rows(records, visual),
    )
    tabulate(
        workbook, 'Validation', 'Input validation log',
        [
            ("level", 'Level'), ("file_name", 'File'), ("sheet_name", 'Worksheet'),
            ("row", 'Row'), ("field", 'Field'), ("value", 'Value'),
            ("message", 'Message'),
        ], issue_rows(issues),
    )
    check_sheet = workbook[visual.text("Validation")]
    for row in range(6, check_sheet.max_row + 1):
        level = check_sheet.cell(row, 1).value
        fill = ERROR if level == "ERROR" else WARNING if level == "WARNING" else None
        if fill:
            for column in range(1, check_sheet.max_column + 1):
                check_sheet.cell(row, column).fill = PatternFill("solid", fgColor=fill)
    profile_sheets = {
        "document_types": ('Document types', 'Document type mentions by dataset'),
        "document_types_by_description": ('Types by inventory', 'Document type mentions by inventory'),
        "document_type_mentions": ('Type mentions', 'Titles and matched document-form expressions'),
        "theme_links": ('Theme links', 'Co-classification network by dataset'),
        "theme_links_by_description": ('Links by inventory', 'Co-classification network by inventory'),
        "theme_link_cases": ('Link evidence', 'Case files contributing theme pairs'),
    }
    for key, (title, heading) in profile_sheets.items():
        tabulate(workbook, title, heading, REPORT_COLUMNS[key], analysis[key],
                       percent_keys={"percent_of_titles"},
                       decimal_keys={"per_1000_titles", "shared_per_1000_titles", "relative_frequency"})
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def _html_table(columns: list[tuple[str, str]], rows: Iterable[dict[str, Any]], limit: int | None = None) -> str:
    selected = list(rows)
    if limit is not None:
        selected = selected[:limit]
    parts = ["<div class='table-wrap'><table><thead><tr>"]
    parts.extend(f"<th>{html.escape(label)}</th>" for _, label in columns)
    parts.append("</tr></thead><tbody>")
    for row in selected:
        parts.append("<tr>")
        for key, _ in columns:
            value = row.get(key, "")
            if isinstance(value, float):
                text = f"{value:.2f}"
            else:
                text = str(value if value is not None else "")
            parts.append(f"<td>{html.escape(text)}</td>")
        parts.append("</tr>")
    parts.append("</tbody></table></div>")
    return "".join(parts)


def write_html_report(
    path: Path,
    datasets: list[Dataset],
    analysis: dict[str, Any],
    chart_files: list[str],
    issues: list[Issue],
    visual: Visualization | None = None,
) -> None:
    visual = visual or Visualization()
    overview = analysis["overview"]
    issue_counts = Counter(issue.level for issue in issues)
    title_label = html.escape(visual.text("Case titles"))
    cards = "".join(
        f"<article class='card'><h3>{html.escape(visual.dataset(row['dataset_id'], row['short_label'], short=True))}</h3>"
        f"<strong>{row['analyzable_titles']:,}</strong><span>{title_label}</span>"
        f"<p>{html.escape(row['reference'])}</p></article>"
        for row in overview
    )
    chart_titles = {
        "dataset_sizes.png": 'Dataset sizes: each bar broken down by inventory',
        "chronology.png": 'Dataset chronology',
        "chronology_by_description.png": 'Annual shares of case files per inventory, %: overlaid series',
        "classification_coverage.png": 'Classification coverage: bars composed of individual inventories',
        "categories.png": 'Thematic categories: adjacent fond panels with a shared legend',
        "document_types.png": 'Document type mentions: share of titles in each dataset, %',
    }
    charts = "".join(
        f"<figure><img src='figures/{html.escape(filename)}' alt='{html.escape(visual.text(chart_titles.get(filename, filename)))}'>"
        f"<figcaption>{html.escape(visual.text(chart_titles.get(filename, filename)))}</figcaption></figure>"
        for filename in chart_files
    )
    top_vocabulary = [row for row in analysis["vocabulary"] if row["type"] == "Distinctive" and row["rank"] <= 15]

    def table(columns, rows, limit=None):
        return _html_table([(key, visual.text(label)) for key, label in columns],
                           visual.rows(rows), limit)

    sections = {
        "CARDS": cards,
        "ERRORS": str(issue_counts["ERROR"]),
        "WARNINGS": str(issue_counts["WARNING"]),
        "INFO": str(issue_counts["INFO"]),
        "CHARTS": charts,
        "TIME": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "OVERVIEW": table([('dataset', 'Dataset'), ('reference', 'Composition'), ('descriptions', 'Inventories'), ('analyzable_titles', 'Case titles'), ('start_year', 'From'), ('end_year', 'To'), ('classification_coverage_percent', 'Classified, %')], overview),
        "COVERAGE": table([('dataset', 'Dataset'), ('analyzable_titles', 'All analyzable titles (100%)'), ('subject_classified', 'Subject classified'), ('subject_classified_percent', 'Classified, %'), ('context_only', 'Context only'), ('context_only_percent', 'Context only, %'), ('unclassified', 'Unclassified'), ('unclassified_percent', 'Unclassified, %')], analysis['classification_coverage']),
        "DESCRIPTIONS": table([('dataset', 'Dataset'), ('reference', 'Archival inventory'), ('analyzable_titles', 'Analyzable titles'), ('classified_titles', 'Subject classified'), ('withdrawn', 'Withdrawn')], analysis['descriptions']),
        "TYPES": table([('dataset', 'Dataset'), ('document_type', 'Document type'), ('titles_total', 'All titles (100%)'), ('titles_with_type', 'Titles with mention'), ('percent_of_titles', 'Share, %'), ('per_1000_titles', 'Per 1,000')], analysis['document_types']),
        "VOCABULARY": table([('dataset', 'Dataset'), ('rank', 'Rank'), ('term', 'Term'), ('titles', 'Case titles'), ('percent_of_titles', 'Share, %')], top_vocabulary),
        "SIMILARITIES": table([('similarity', 'Similarity'), ('reference_a', 'Reference A'), ('case_a', 'Case A'), ('title_a', 'Title A'), ('reference_b', 'Reference B'), ('case_b', 'Case B'), ('title_b', 'Title B')], analysis['similarities'], 100),
    }
    source = (Path(__file__).parent / "templates" / "report.html").read_text(encoding="utf-8")
    template = visual.template(source)
    page = re.sub(r"__([A-Z]+)__", lambda match: sections[match[1]], template)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(page, encoding="utf-8")


def write_manifest(
    path: Path,
    base_dir: Path,
    datasets: list[Dataset],
    dictionary_version: str,
    output_files: list[Path],
    visual: Visualization | None = None,
) -> None:
    visual = visual or Visualization()
    inputs = [dataset.path for dataset in datasets]
    payload = {
        "status": "complete",
        "script_version": __version__,
        "dictionary_version": dictionary_version,
        "document_type_rules_version": DOCUMENT_TYPE_RULES_VERSION,
        "visualization": visual.metadata(),
        "presentation_sha256": {
            name: hashlib.sha256((base_dir / name).read_bytes()).hexdigest()
            for name in ("src/localization.py", "config/localization.json", "src/charts.py",
                         "src/templates/report.html",
                         "config/analysis.yaml", "src/templates/theme_links.html",
                         "src/reports.py", "src/comparative_profiles.py")
            if (base_dir / name).is_file()
        },
        "completed_at": datetime.now().isoformat(timespec="seconds"),
        "python": platform.python_version(),
        "dependencies": {
            name: version(name) for name in ("openpyxl", "PyYAML", "matplotlib", "scikit-learn")
        },
        "analysis_sha256": {
            str(item.relative_to(base_dir)): hashlib.sha256(item.read_bytes()).hexdigest()
            for item in sorted([
                *base_dir.glob("src/*.py"), *base_dir.glob("config/*.yaml"),
                *base_dir.glob("dictionaries/*.yaml"), *base_dir.glob("main.py"),
                *base_dir.glob("requirements.txt"),
            ])
        },
        "datasets": [
            {"id": dataset.id, "label": dataset.label, "file": dataset.path.name, "reference": dataset.reference}
            for dataset in datasets
        ],
        "sha256": {
            item.name: hashlib.sha256(item.read_bytes()).hexdigest()
            for item in inputs
        },
        "outputs": [str(item.relative_to(path.parent)) for item in output_files if item.exists()],
    }
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
