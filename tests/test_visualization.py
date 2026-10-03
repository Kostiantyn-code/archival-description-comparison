from __future__ import annotations

import argparse
from copy import deepcopy
import io
import json
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
import unittest
from contextlib import redirect_stdout
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
import yaml

import main
from src.analysis import build_analysis
from src.classification import classify_records, load_language_dictionaries
from src.comparative_profiles import write_theme_network
from src.document_types import DOCUMENT_TYPES
from src.loader import load_all
from src.reports import _description_label
from src.visualization import Visualization, english_catalog


BASE = Path(__file__).resolve().parents[1]
OVERRIDES = {
    "datasets": {"left": "Combined collection", "right": "Second collection", "empty": "Empty collection"},
    "dataset_short": {"left": "Collection A", "right": "Collection B", "empty": "Empty"},
    "archives": {"Архів А": "Archive A", "Архів Б": "Archive B"},
}


def make_inputs(root: Path) -> dict:
    """Same sheet names and case numbers, mixed archives, Russian titles, empty dataset."""
    datasets = []
    for key, inventory_count in (("left", 2), ("right", 1), ("empty", 1)):
        book = Workbook()
        book.remove(book.active)
        for number in range(1, inventory_count + 1):
            sheet = book.create_sheet(f"Опис {number}")
            sheet.append(["Архів", "Архів Б" if number == 2 else "Архів А"])
            sheet.append(["Фонд", "230" if number == 1 else "356"])
            sheet.append(["Опис", str(number)])
            sheet.append(["Мова", "ru" if number == 2 else "uk"])
            sheet.append(["№ справи", "Заголовок справи", "Крайні дати", "Кількість аркушів", "Примітки"])
            if key != "empty":
                titles = (["Переписка о школе и больнице", "Рапорты о школе"] if number == 2 else [
                    "Листування про відкриття школи та лікарні",
                    "Рапорти про міську лікарню", "Заголовок без теми",
                    '</script><img src=x onerror="alert(1)">',
                ])
                for row, title in enumerate(titles):
                    sheet.append([1, title, str(1850 + row), 10, None])
        book.save(root / f"{key}.xlsx")
        book.close()
        datasets.append({"pattern": f"{key}.xlsx", "id": key,
                         "label": f"Масив {key}", "short_label": f"Масив {key}"})
    return {"datasets": datasets}


def payload(page: str) -> dict:
    return json.loads(page.split('id="tl-data">')[1].split('</script>')[0])


def workbook_cells(path: Path) -> dict:
    book = load_workbook(path, read_only=True)
    try:
        return {sheet.title: list(sheet.values) for sheet in book}
    finally:
        book.close()


class VisualizationTests(unittest.TestCase):
    def test_defaults_validation_overrides_and_original_fallback(self):
        self.assertEqual(Visualization.from_config({}).metadata(),
                         {"language": "uk", "bilingual_labels": False, "label_overrides": {}})
        invalid = [None, [], {"language": "de"}, {"bilingual_labels": "false"},
                   {"label_overrides": []}, {"label_overrides": {"datasets": {"a": 1}}}]
        for options in invalid:
            with self.subTest(options=options), self.assertRaisesRegex(ValueError, "visualization"):
                Visualization.from_config({"visualization": options})
        visual = Visualization("en", True, OVERRIDES)
        self.assertEqual(visual.dataset("left", "Масив А", short=True), "Collection A (Масив А)")
        self.assertEqual(Visualization("en", False, {"datasets": {"a": "A"}}).dataset("a", "А", short=True), "A")
        self.assertEqual(visual.label("categories", "custom", "Нова тема"), "Нова тема")
        self.assertEqual(visual.inventory_name("Опис 12"), "Inventory 12 (Опис 12)")
        self.assertEqual(visual.inventory_name("Авторська назва"), "Авторська назва")
        self.assertEqual(visual.inventory_reference("a:Опис 1", "", "", "", "Опис 1"), "Inventory 1 (Опис 1)")
        self.assertEqual(visual.inventory_reference("a:Назва", "", "", "", "Назва"), "Назва")
        self.assertEqual(Visualization("uk", True, OVERRIDES).dataset("left", "Масив А"), "Масив А")

    def test_catalog_covers_current_taxonomy_and_document_types(self):
        categories, _, blocks, _ = load_language_dictionaries(BASE, yaml, "uk")
        catalog = english_catalog()
        for group, keys in (("categories", [c.id for c in categories]),
                            ("category_short", [c.id for c in categories]),
                            ("macroblocks", blocks), ("macroblock_short", blocks),
                            ("document_types", [kind.id for kind in DOCUMENT_TYPES])):
            self.assertTrue(set(keys) <= set(catalog[group]), group)
            self.assertTrue(all(not re.search("[А-Яа-яІіЇїЄєҐґ]", catalog[group][key]) for key in keys))

    def test_template_localization_is_complete_and_context_safe(self):
        source = (BASE / "src/templates/theme_links.html").read_text(encoding="utf-8")
        self.assertEqual(Visualization().template(source), source)
        translated = Visualization("en").template(source)
        self.assertIn('lang="en"', translated)
        self.assertIn("en-GB", translated)
        self.assertIsNone(re.search("[А-Яа-яІіЇїЄєҐґ]", translated))
        # Translation strings containing quotes/markup must not break JS or HTML.
        hostile = '</script>\'` ${alert(1)} & "'
        with patch("src.visualization.english_catalog", return_value={"ui": {"Масив": hostile}}):
            page = Visualization("en").template('<p>Масив</p><script>const x="Масив";</script>')
        self.assertIn("&lt;/script&gt;", page)
        self.assertEqual(page.count("</script>"), 1)
        self.assertIn(r"\u0024", page)

    @unittest.skipUnless(shutil.which("node"), "Node.js is optional; needed for generated JavaScript syntax check")
    def test_localized_network_javascript_parses(self):
        source = (BASE / "src/templates/theme_links.html").read_text(encoding="utf-8")
        for visual in (Visualization(), Visualization("en"), Visualization("en", True)):
            script = re.findall(r"<script>(.*?)</script>", visual.template(source), re.S)[0]
            result = subprocess.run([shutil.which("node"), "--check"], input=script,
                                    text=True, capture_output=True)
            self.assertEqual(result.returncode, 0, result.stderr)

    def test_network_preserves_counts_evidence_ids_and_source_objects(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            config = make_inputs(root)
            datasets, descriptions, records, _ = load_all(root, config, {})
            categories, ambiguities, blocks, _ = load_language_dictionaries(BASE, yaml, "uk")
            russian, ru_ambiguities, _, _ = load_language_dictionaries(BASE, yaml, "ru")
            classify_records(records, {"uk": categories, "ru": russian},
                             {"uk": ambiguities, "ru": ru_ambiguities}, list(blocks))
            analysis = build_analysis(BASE, datasets, descriptions, records, categories,
                                      {"similarity": {"enabled": False}}, blocks)
            original = deepcopy((analysis, datasets, descriptions, records, categories))
            network = analysis["theme_network"]
            for visual in (Visualization(), Visualization("en", False, OVERRIDES),
                           Visualization("en", True, OVERRIDES)):
                path = root / "theme_links.html"
                write_theme_network(path, network, visual, datasets, descriptions)
                page = path.read_text(encoding="utf-8")
                displayed = payload(page)
                self.assertEqual(displayed["scopes"], network["scopes"])
                self.assertEqual(displayed["meta"]["catOrder"], network["meta"]["catOrder"])
                self.assertEqual(displayed["meta"]["catBlock"], network["meta"]["catBlock"])
                for field in ("scopes", "datasets"):
                    self.assertEqual([key for key, _ in displayed["meta"][field]],
                                     [key for key, _ in network["meta"][field]])
                if visual.language == "uk":
                    self.assertEqual(displayed, network)
                else:
                    self.assertIn("Inventory 1", displayed["meta"]["scopes"][1][1])
                    self.assertIn("Archive A", displayed["meta"]["scopes"][1][1])
                self.assertNotIn('</script><img src=x', page)
                self.assertEqual((analysis, datasets, descriptions, records, categories), original)
            visual = Visualization("en", True, {"descriptions": {"left:Опис 1": "Custom reference"}})
            item = next(d for d in descriptions if d.dataset_id == "left" and d.sheet_name == "Опис 1")
            self.assertIn("Custom reference (", _description_label(item, visual))
            write_theme_network(path, network, visual, datasets, descriptions)
            self.assertTrue(any("Custom reference (" in label
                                for _, label in payload(path.read_text())["meta"]["scopes"]))

    def test_main_three_languages_preserve_csv_and_workbook_and_export_six_charts(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            inputs = root / "input"
            inputs.mkdir()
            dataset_config = make_inputs(inputs)
            base_config = yaml.safe_load((BASE / "config/analysis.yaml").read_text(encoding="utf-8"))
            snapshots = []
            titles = {
                "dataset_sizes": "Обсяг порівнюваних масивів за описами",
                "chronology": "Хронологічний розподіл справ",
                "chronology_by_description": "Хронологія описів у кожному порівнюваному масиві",
                "classification_coverage": "Покриття класифікацією за окремими описами",
                "categories": "Тематичні категорії за описами",
                "document_types": "Згадки типів документів, % заголовків масиву",
            }
            for mode, options in (("uk", None),
                                  ("en", {"language": "en", "label_overrides": OVERRIDES}),
                                  ("bilingual", {"language": "en", "bilingual_labels": True, "label_overrides": OVERRIDES})):
                config = deepcopy(base_config)
                config.pop("visualization")
                if options is not None:
                    config["visualization"] = options
                args = argparse.Namespace(input_dir=str(inputs), output_dir=str(root / "output"),
                                          run_name=mode, no_similarity=False)
                with patch("main.parse_args", return_value=args), \
                     patch("main.safe_load_yaml", side_effect=[config, dataset_config]), \
                     redirect_stdout(io.StringIO()):
                    output = main.main()
                snapshots.append((
                    {path.name: path.read_bytes() for path in (output / "tables").glob("*.csv")},
                    workbook_cells(output / "comparison.xlsx"),
                ))
                manifest = json.loads((output / "run_manifest.json").read_text())
                self.assertEqual(manifest["visualization"], Visualization.from_config(config).metadata())
                self.assertIn("config/visualization.en.json", manifest["presentation_sha256"])
                self.assertEqual(len(list((output / "figures").glob("*.png"))), 6)
                self.assertEqual(len(list((output / "figures").glob("*.svg"))), 6)
                self.assertEqual(sum(name.endswith(".svg") for name in manifest["outputs"]), 6)
                visual = Visualization.from_config(config)
                for stem, title in titles.items():
                    svg = (output / "figures" / f"{stem}.svg").read_text(encoding="utf-8")
                    self.assertIn(visual.text(title), svg)
                    if mode == "en":
                        self.assertIsNone(re.search("[А-Яа-яІіЇїЄєҐґ]", svg), stem)
                categories_svg = (output / "figures/categories.svg").read_text()
                self.assertIn("Education" if mode != "uk" else "Освіта", categories_svg)
                if mode == "bilingual":
                    self.assertIn("Education (Освіта)", categories_svg)
                html = (output / "report.html").read_text()
                self.assertEqual(html.count("<img "), 6)
                if mode != "uk":
                    self.assertIn("Open the interactive co-classification network", html)
                    self.assertIn("Листування про відкриття школи", html)
            self.assertEqual(snapshots[0], snapshots[1])
            self.assertEqual(snapshots[0], snapshots[2])


if __name__ == "__main__":
    unittest.main()
